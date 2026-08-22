"""Export search-result URLs to an Excel workbook for human relevance review.

Accuracy is the only metric that needs human judgement, so every unique URL a
method returned has to be opened and labeled once. This script collects the
result JSON files of one experiment, removes duplicates, and writes a workbook
with one row per URL and a true/false dropdown.

Two design decisions matter for the validity of the labels:

Predictions are not shown.
    The workbook deliberately omits ``predicted_relevant`` and
    ``relevance_score``. Showing a method's guess to the person labeling the
    data biases the ground truth toward that method and inflates its accuracy.

One shared label per URL.
    A URL is labeled once, not once per method. Relevance is a property of the
    page and the keyword, not of the method that happened to find it. This
    matches the ``method``-blank convention of the existing ground-truth file
    and keeps the four methods comparable on identical labels.

Fill in ``human_relevant`` (and optionally ``human_note``) for every row, then
feed the file back through ``import_label_workbook.py``.
"""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
from typing import Any


try:
    from .ground_truth import GroundTruthError, normalize_url
    from .result_validator import find_result_files, validate_result_file
except ImportError:
    from ground_truth import GroundTruthError, normalize_url
    from result_validator import find_result_files, validate_result_file


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUTPUT_PATH = (
    PROJECT_ROOT
    / "data"
    / "labels"
    / "label_review_workbook.xlsx"
)

REVIEW_COLUMNS = (
    "row_number",
    "product_id",
    "category_group",
    "keyword",
    "domain",
    "url",
    "found_by_methods",
    "human_relevant",
    "human_note",
    "reviewer",
)

RELEVANCE_RULE = (
    "A result is RELEVANT when both conditions hold:\n"
    "  1. It is the product named in the keyword, including the variant "
    "(storage, volume, weight, size, edition) when the keyword states one.\n"
    "  2. The page has transactional purpose: a retailer product page, a "
    "marketplace listing, a classified listing, or a price-comparison page.\n"
    "\n"
    "A result is IRRELEVANT when:\n"
    "  - it is a different product or a different variant,\n"
    "  - it is an accessory for the product rather than the product,\n"
    "  - it is a news article, a blog post, a review, or a forum thread,\n"
    "  - it is a category or search page that does not reach the product.\n"
    "\n"
    "Being out of stock does NOT make a page irrelevant.\n"
    "If you are unsure, write your reasoning in human_note and still choose "
    "the value the rule above implies."
)


def load_category_groups_by_product(
    subset_file: Path | None,
) -> dict[str, str]:
    """Map product_id to category_group so the reviewer sees the category."""

    if subset_file is None or not subset_file.exists():
        return {}

    # The curated subset is a CSV; the processor writes JSON. Both already
    # carry product_id and category_group, so a review workbook can be built
    # from whichever file the run in question produced.
    if subset_file.suffix.lower() == ".json":
        records = json.loads(subset_file.read_text(encoding="utf-8"))

        if not isinstance(records, list):
            raise GroundTruthError(
                f"Expected a list of products in {subset_file}"
            )

        rows: list[dict] = records
    else:
        with subset_file.open(
            mode="r",
            encoding="utf-8-sig",
            newline="",
        ) as stream:
            rows = list(csv.DictReader(stream))

    return {
        str(row["product_id"]).strip(): str(row.get("category_group") or "").strip()
        for row in rows
        if str(row.get("product_id", "")).strip()
    }


def review_url_key(url: str) -> str:
    """Collapse the URL spellings that point at the same page.

    Deduplication here has to agree with ``ground_truth.normalize_url``. When
    it does not, a trailing slash is enough to send the same page to a
    reviewer twice and then collide when the two labels are loaded under one
    normalized key.
    """

    try:
        return normalize_url(url)
    except GroundTruthError:
        # A result URL that cannot be normalized is still worth reviewing;
        # fall back to the raw spelling instead of dropping the row.
        return url.strip().casefold()


def load_already_reviewed(paths: list[Path]) -> set[tuple[str, str]]:
    """Read (product_id, url) pairs out of workbooks that are already filled.

    A provider daily limit can split one experiment across sessions. Reviewers
    should never be handed a URL they already judged, so the second workbook
    excludes everything the first one covered.
    """

    try:
        from openpyxl import load_workbook
    except ImportError as error:  # pragma: no cover - dependency guard
        raise RuntimeError(
            "openpyxl is required to read an existing review workbook."
        ) from error

    reviewed: set[tuple[str, str]] = set()

    for path in paths:
        if not path.exists():
            raise FileNotFoundError(f"Workbook to exclude was not found: {path}")

        workbook = load_workbook(path, data_only=True)

        if "URLs" not in workbook.sheetnames:
            continue

        sheet = workbook["URLs"]
        rows = sheet.iter_rows(values_only=True)

        try:
            header = [
                str(value).strip() if value is not None else ""
                for value in next(rows)
            ]
        except StopIteration:
            continue

        if "product_id" not in header or "url" not in header:
            continue

        product_index = header.index("product_id")
        url_index = header.index("url")

        for values in rows:
            if product_index >= len(values) or url_index >= len(values):
                continue

            product_id = str(values[product_index] or "").strip()
            url = str(values[url_index] or "").strip()

            if product_id and url:
                reviewed.add((product_id.casefold(), review_url_key(url)))

    return reviewed


def collect_review_rows(
    results_directory: Path,
    category_by_product: dict[str, str] | None = None,
    already_reviewed: set[tuple[str, str]] | None = None,
) -> list[dict[str, Any]]:
    """Read result files and return one deduplicated row per URL."""

    category_by_product = category_by_product or {}
    already_reviewed = already_reviewed or set()
    rows: dict[tuple[str, str, str], dict[str, Any]] = {}

    for file_path in find_result_files(results_directory):
        payload = validate_result_file(file_path)
        product_id = str(payload["product_id"]).strip()
        keyword = str(payload["keyword"]).strip()
        method = str(payload["method"]).strip()

        for result in payload["results"]:
            url = str(result["url"]).strip()

            if (product_id.casefold(), review_url_key(url)) in already_reviewed:
                continue

            key = (product_id, keyword, review_url_key(url))
            existing = rows.get(key)

            if existing is None:
                rows[key] = {
                    "product_id": product_id,
                    "category_group": category_by_product.get(product_id, ""),
                    "keyword": keyword,
                    "domain": str(result["domain"]).strip(),
                    "url": url,
                    "methods": {method},
                }
            else:
                existing["methods"].add(method)

    ordered = sorted(
        rows.values(),
        key=lambda row: (row["product_id"], row["domain"], row["url"]),
    )

    for row_number, row in enumerate(ordered, start=1):
        row["row_number"] = row_number
        row["found_by_methods"] = ", ".join(sorted(row.pop("methods")))
        row["human_relevant"] = ""
        row["human_note"] = ""
        row["reviewer"] = ""

    return ordered


def write_workbook(
    rows: list[dict[str, Any]],
    output_path: Path,
) -> Path:
    """Write the review workbook with an instructions sheet and a dropdown."""

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError as error:  # pragma: no cover - dependency guard
        raise RuntimeError(
            "openpyxl is required to export the review workbook. "
            "Install it with: pip install -r requirements.txt"
        ) from error

    workbook = Workbook()

    instructions = workbook.active
    instructions.title = "Nasıl doldurulur"
    instructions["A1"] = "Relevance review instructions"
    instructions["A1"].font = Font(bold=True, size=14)
    instructions["A3"] = RELEVANCE_RULE
    instructions["A3"].alignment = Alignment(
        wrap_text=True,
        vertical="top",
    )
    instructions.merge_cells("A3:H30")
    instructions.column_dimensions["A"].width = 110
    instructions["A32"] = (
        "Open every URL on the 'URLs' sheet, choose true or false in the "
        "human_relevant column, and write your name in reviewer. "
        "Leave no row blank."
    )
    instructions["A32"].alignment = Alignment(wrap_text=True)

    sheet = workbook.create_sheet("URLs")
    sheet.append(list(REVIEW_COLUMNS))

    header_fill = PatternFill("solid", fgColor="DDEBF7")
    answer_fill = PatternFill("solid", fgColor="FFF2CC")

    for column_index, _ in enumerate(REVIEW_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=column_index)
        cell.font = Font(bold=True)
        cell.fill = header_fill

    for row in rows:
        sheet.append([row.get(column, "") for column in REVIEW_COLUMNS])

    url_column = REVIEW_COLUMNS.index("url") + 1
    answer_column = REVIEW_COLUMNS.index("human_relevant") + 1
    note_column = REVIEW_COLUMNS.index("human_note") + 1
    reviewer_column = REVIEW_COLUMNS.index("reviewer") + 1

    for row_index in range(2, len(rows) + 2):
        url_cell = sheet.cell(row=row_index, column=url_column)
        url_cell.hyperlink = url_cell.value
        url_cell.font = Font(color="0563C1", underline="single")

        for column_index in (answer_column, note_column, reviewer_column):
            sheet.cell(row=row_index, column=column_index).fill = answer_fill

    validation = DataValidation(
        type="list",
        formula1='"true,false"',
        allow_blank=True,
    )
    validation.error = "Choose true or false."
    validation.errorTitle = "Invalid relevance label"
    sheet.add_data_validation(validation)
    answer_letter = get_column_letter(answer_column)
    validation.add(f"{answer_letter}2:{answer_letter}{len(rows) + 1}")

    widths = {
        "row_number": 10,
        "product_id": 12,
        "category_group": 26,
        "keyword": 42,
        "domain": 24,
        "url": 70,
        "found_by_methods": 34,
        "human_relevant": 16,
        "human_note": 46,
        "reviewer": 14,
    }

    for column_index, column_name in enumerate(REVIEW_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(column_index)].width = (
            widths.get(column_name, 18)
        )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(REVIEW_COLUMNS))}{len(rows) + 1}"
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Export unique search-result URLs to an Excel review workbook."
        )
    )
    parser.add_argument(
        "results",
        type=Path,
        help="Directory holding the experiment result JSON files.",
    )
    parser.add_argument(
        "--subset",
        type=Path,
        default=(
            PROJECT_ROOT
            / "data"
            / "evaluation"
            / "evaluation_subset_multicategory.csv"
        ),
        help="Evaluation subset used to label each row with its category.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT_PATH,
    )
    parser.add_argument(
        "--exclude",
        type=Path,
        nargs="*",
        default=[],
        help=(
            "Workbooks whose URLs were already reviewed. Their rows are left "
            "out, so a second session only receives new URLs."
        ),
    )
    args = parser.parse_args()

    already_reviewed = load_already_reviewed(list(args.exclude))
    rows = collect_review_rows(
        results_directory=args.results,
        category_by_product=load_category_groups_by_product(args.subset),
        already_reviewed=already_reviewed,
    )

    if not rows:
        print(
            "[ERROR] No new URLs were found. Every result is already covered "
            "by the excluded workbook(s)."
        )
        raise SystemExit(1)

    output_path = write_workbook(rows, args.output)

    product_count = len({row["product_id"] for row in rows})
    category_count = len(
        {row["category_group"] for row in rows if row["category_group"]}
    )

    print("=== LABEL REVIEW WORKBOOK ===")
    print(f"URLs to review:  {len(rows)}")

    if already_reviewed:
        print(f"Already reviewed: {len(already_reviewed)} (excluded)")

    print(f"Products:        {product_count}")
    print(f"Category groups: {category_count}")
    print(f"Workbook:        {output_path}")


if __name__ == "__main__":
    main()
