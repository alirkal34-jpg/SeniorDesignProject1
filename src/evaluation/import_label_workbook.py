"""Read a completed relevance-review workbook back into a ground-truth CSV.

This is the return leg of ``export_label_workbook.py``. It validates that the
reviewers filled every row, normalizes the true/false answers, and writes the
``product_id,keyword,method,domain,url,human_relevant,notes`` file the metrics
modules already understand.

The ``method`` column is left blank on purpose: each URL carries one shared
label that applies to every method that returned it, so the four methods are
always scored against identical ground truth.

The import refuses to write a partial file. An unlabeled or unreadable row is
reported with its row number so it can be fixed in the workbook instead of
silently dropping out of the accuracy calculation.

A reviewer may leave a row unresolved on purpose, for example by writing
``true?`` when a listing sells the right product in a bundle. Those rows are
settled in a separate adjudication file rather than by editing the workbook, so
the human record stays exactly as the reviewer left it and every adjudicated
label can be traced back to the rule that produced it. An adjudication only
applies to a row the workbook itself leaves unreadable: it can never overturn a
true or false the reviewer already wrote, and an entry that matches no such row
fails the import instead of passing unnoticed.

URLs are compared the way the metrics compare them, so two spellings of one
page — most often a trailing slash — collapse into a single record when the
reviewer answered both the same way, and stop the import when the two answers
disagree.
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path
from typing import Any


try:
    from .export_label_workbook import review_url_key
except ImportError:
    from export_label_workbook import review_url_key


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_WORKBOOK_PATH = (
    PROJECT_ROOT
    / "data"
    / "labels"
    / "label_review_workbook.xlsx"
)
DEFAULT_OUTPUT_PATH = (
    PROJECT_ROOT
    / "data"
    / "labels"
    / "multicategory_ground_truth.csv"
)

GROUND_TRUTH_COLUMNS = (
    "product_id",
    "keyword",
    "method",
    "domain",
    "url",
    "human_relevant",
    "notes",
)

TRUE_VALUES = {"true", "1", "yes", "y", "evet", "dogru", "doğru"}
FALSE_VALUES = {"false", "0", "no", "n", "hayir", "hayır", "yanlis", "yanlış"}

REQUIRED_WORKBOOK_COLUMNS = (
    "product_id",
    "keyword",
    "domain",
    "url",
    "human_relevant",
)

REQUIRED_ADJUDICATION_COLUMNS = (
    "product_id",
    "url",
    "resolved_label",
    "rule",
    "adjudicated_by",
)

AdjudicationKey = tuple[str, str]


class LabelImportError(ValueError):
    """Raised when the completed workbook cannot be used as ground truth."""


def read_workbook_rows(
    workbook_path: Path,
    sheet_name: str = "URLs",
) -> list[dict[str, Any]]:
    """Read the review sheet into dictionaries keyed by column name."""

    try:
        from openpyxl import load_workbook
    except ImportError as error:  # pragma: no cover - dependency guard
        raise LabelImportError(
            "openpyxl is required to read the review workbook. "
            "Install it with: pip install -r requirements.txt"
        ) from error

    if not workbook_path.exists():
        raise LabelImportError(
            f"Review workbook was not found: {workbook_path}"
        )

    workbook = load_workbook(workbook_path, data_only=True)

    if sheet_name not in workbook.sheetnames:
        raise LabelImportError(
            f"Workbook has no '{sheet_name}' sheet. "
            f"Available sheets: {workbook.sheetnames}"
        )

    sheet = workbook[sheet_name]
    row_iterator = sheet.iter_rows(values_only=True)

    try:
        header = next(row_iterator)
    except StopIteration as error:
        raise LabelImportError("Review sheet is empty.") from error

    column_names = [
        str(value).strip() if value is not None else ""
        for value in header
    ]
    missing_columns = [
        column
        for column in REQUIRED_WORKBOOK_COLUMNS
        if column not in column_names
    ]

    if missing_columns:
        raise LabelImportError(
            f"Review sheet is missing columns: {missing_columns}"
        )

    rows: list[dict[str, Any]] = []

    for sheet_row_number, values in enumerate(row_iterator, start=2):
        if all(value is None or str(value).strip() == "" for value in values):
            continue

        row = {
            column_name: values[index] if index < len(values) else None
            for index, column_name in enumerate(column_names)
            if column_name
        }
        row["_sheet_row"] = sheet_row_number
        rows.append(row)

    return rows


def parse_label(value: Any) -> bool | None:
    """Normalize one human answer, returning None when it is unusable."""

    if value is None:
        return None

    if isinstance(value, bool):
        return value

    text = str(value).strip().casefold()

    if text in TRUE_VALUES:
        return True

    if text in FALSE_VALUES:
        return False

    return None


def adjudication_key(product_id: str, url: str) -> AdjudicationKey:
    """Build the key that ties an adjudication to one workbook row."""

    return (product_id.strip().casefold(), url.strip().casefold())


def load_adjudications(
    adjudication_path: Path,
) -> dict[AdjudicationKey, dict[str, str]]:
    """Read the decisions that settle rows the reviewer left unresolved."""

    if not adjudication_path.exists():
        raise LabelImportError(
            f"Adjudication file was not found: {adjudication_path}"
        )

    with adjudication_path.open(encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        field_names = reader.fieldnames or []
        missing_columns = [
            column
            for column in REQUIRED_ADJUDICATION_COLUMNS
            if column not in field_names
        ]

        if missing_columns:
            raise LabelImportError(
                f"Adjudication file is missing columns: {missing_columns}"
            )

        adjudications: dict[AdjudicationKey, dict[str, str]] = {}
        problems: list[str] = []

        for line_number, row in enumerate(reader, start=2):
            product_id = str(row.get("product_id") or "").strip()
            url = str(row.get("url") or "").strip()
            rule = str(row.get("rule") or "").strip()
            adjudicated_by = str(row.get("adjudicated_by") or "").strip()
            label = parse_label(row.get("resolved_label"))

            if not (product_id and url):
                problems.append(
                    f"Line {line_number}: product_id and url are required."
                )
                continue

            if label is None:
                problems.append(
                    f"Line {line_number} ({product_id}): resolved_label is "
                    f"missing or unreadable "
                    f"(found {row.get('resolved_label')!r}). Use true "
                    "or false."
                )
                continue

            if not (rule and adjudicated_by):
                problems.append(
                    f"Line {line_number} ({product_id}): rule and "
                    "adjudicated_by are required so the decision stays "
                    "traceable."
                )
                continue

            key = adjudication_key(product_id, url)

            if key in adjudications:
                problems.append(
                    f"Line {line_number} ({product_id}): duplicate "
                    "adjudication for this URL."
                )
                continue

            adjudications[key] = {
                "product_id": product_id,
                "url": url,
                "label": "true" if label else "false",
                "rule": rule,
                "adjudicated_by": adjudicated_by,
            }

    if problems:
        raise LabelImportError(
            f"{len(problems)} adjudication row(s) need attention:\n  "
            + "\n  ".join(problems)
        )

    if not adjudications:
        raise LabelImportError(
            f"Adjudication file has no decisions: {adjudication_path}"
        )

    return adjudications


def build_ground_truth_records(
    rows: list[dict[str, Any]],
    adjudications: dict[AdjudicationKey, dict[str, str]] | None = None,
) -> list[dict[str, str]]:
    """Convert reviewed rows into ground-truth records, or fail loudly."""

    records: list[dict[str, str]] = []
    problems: list[str] = []
    seen_keys: dict[tuple[str, str, str], tuple[bool, str]] = {}
    adjudications = adjudications or {}
    used_adjudications: set[AdjudicationKey] = set()

    for row in rows:
        workbook_name = row.get("_workbook")
        sheet_row = (
            f"{workbook_name} row {row.get('_sheet_row', '?')}"
            if workbook_name
            else f"Row {row.get('_sheet_row', '?')}"
        )
        product_id = str(row.get("product_id") or "").strip()
        keyword = str(row.get("keyword") or "").strip()
        domain = str(row.get("domain") or "").strip()
        url = str(row.get("url") or "").strip()

        if not (product_id and keyword and domain and url):
            problems.append(
                f"{sheet_row}: product_id, keyword, domain and url are "
                "all required."
            )
            continue

        raw_value = row.get("human_relevant")
        label = parse_label(raw_value)
        adjudication = adjudications.get(adjudication_key(product_id, url))
        adjudication_note = ""

        if adjudication is not None and label is not None:
            used_adjudications.add(adjudication_key(product_id, url))
            problems.append(
                f"{sheet_row} ({product_id}, {domain}): the reviewer "
                f"answered {raw_value!r}, so this row must not be "
                "adjudicated. Remove it from the adjudication file."
            )
            continue

        if label is None and adjudication is not None:
            used_adjudications.add(adjudication_key(product_id, url))
            label = adjudication["label"] == "true"
            adjudication_note = (
                f"[adjudicated: reviewer wrote {str(raw_value).strip()!r}; "
                f"resolved to {adjudication['label']} by "
                f"{adjudication['rule']}; "
                f"by: {adjudication['adjudicated_by']}]"
            )

        if label is None:
            problems.append(
                f"{sheet_row} ({product_id}, {domain}): "
                f"human_relevant is missing or unreadable "
                f"(found {raw_value!r}). Use true or false."
            )
            continue

        key = (
            product_id.casefold(),
            keyword.casefold(),
            review_url_key(url),
        )
        seen = seen_keys.get(key)

        if seen is not None:
            seen_label, seen_row = seen

            if seen_label != label:
                problems.append(
                    f"{sheet_row}: {product_id} has contradictory labels "
                    f"for the same page ({seen_row} says "
                    f"{'true' if seen_label else 'false'}, this row says "
                    f"{'true' if label else 'false'}). Decide one answer."
                )

            # The same page can reach the workbook under two spellings, for
            # example with and without a trailing slash. Agreeing labels
            # collapse into the single record the metrics expect.
            continue

        seen_keys[key] = (label, sheet_row)

        note = str(row.get("human_note") or "").strip()
        reviewer = str(row.get("reviewer") or "").strip()

        if reviewer:
            note = f"{note} [reviewer: {reviewer}]".strip()

        if adjudication_note:
            note = f"{note} {adjudication_note}".strip()

        records.append(
            {
                "product_id": product_id,
                "keyword": keyword,
                # Blank means the label is shared by every method.
                "method": "",
                "domain": domain,
                "url": url,
                "human_relevant": "true" if label else "false",
                "notes": note,
            }
        )

    for key in sorted(set(adjudications) - used_adjudications):
        entry = adjudications[key]
        problems.append(
            f"Adjudication for {entry['product_id']} ({entry['url']}) "
            "matched no workbook row. Remove the stale entry or correct "
            "the URL."
        )

    if problems:
        raise LabelImportError(
            f"{len(problems)} row(s) need attention before import:\n  "
            + "\n  ".join(problems)
        )

    if not records:
        raise LabelImportError("No labeled rows were found in the workbook.")

    return records


def write_ground_truth(
    records: list[dict[str, str]],
    output_path: Path,
) -> Path:
    """Write the ground-truth CSV used by the metrics modules."""

    output_path.parent.mkdir(parents=True, exist_ok=True)

    with output_path.open(
        mode="w",
        encoding="utf-8",
        newline="",
    ) as stream:
        writer = csv.DictWriter(
            stream,
            fieldnames=list(GROUND_TRUTH_COLUMNS),
        )
        writer.writeheader()
        writer.writerows(records)

    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Import a completed relevance-review workbook as ground truth."
        )
    )
    parser.add_argument(
        "--workbook",
        type=Path,
        nargs="+",
        default=[DEFAULT_WORKBOOK_PATH],
        help=(
            "One or more completed workbooks. Several are accepted because a "
            "provider daily limit can split one experiment across sessions."
        ),
    )
    parser.add_argument(
        "--sheet",
        default="URLs",
    )
    parser.add_argument(
        "--adjudication",
        type=Path,
        default=None,
        help=(
            "CSV holding the decisions for rows the reviewer left "
            "unresolved. Columns: "
            f"{', '.join(REQUIRED_ADJUDICATION_COLUMNS)}."
        ),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT_PATH,
    )
    args = parser.parse_args()

    try:
        rows: list[dict[str, Any]] = []

        for workbook_path in args.workbook:
            workbook_rows = read_workbook_rows(workbook_path, args.sheet)

            for row in workbook_rows:
                row["_workbook"] = workbook_path.name

            rows.extend(workbook_rows)

        adjudications = (
            load_adjudications(args.adjudication)
            if args.adjudication is not None
            else None
        )
        records = build_ground_truth_records(rows, adjudications)
    except LabelImportError as error:
        print(f"[ERROR] {error}")
        raise SystemExit(1) from error

    output_path = write_ground_truth(records, args.output)

    relevant_count = sum(
        1
        for record in records
        if record["human_relevant"] == "true"
    )

    collapsed_count = len(rows) - len(records)
    adjudicated_count = sum(
        1
        for record in records
        if "[adjudicated:" in record["notes"]
    )

    print("=== GROUND TRUTH IMPORTED ===")
    print(f"Labeled URLs:   {len(records)}")
    print(f"Relevant:       {relevant_count}")
    print(f"Irrelevant:     {len(records) - relevant_count}")
    print(f"Adjudicated:    {adjudicated_count}")
    print(f"Merged repeats: {collapsed_count}")
    print(f"Products:       {len({r['product_id'] for r in records})}")
    print(f"Ground truth:   {output_path}")


if __name__ == "__main__":
    main()
